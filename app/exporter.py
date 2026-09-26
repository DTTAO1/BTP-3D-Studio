from __future__ import annotations

from pathlib import Path
from typing import Any
import csv
import html
import json
import math

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


CATEGORY_COLOR = {
    "foundation": "F2B84B",
    "slab": "4CA7E8",
    "column": "A778E6",
    "beam": "EA746A",
    "wall": "35C597",
    "fill": "9B8A5C",
    "opening": "D7DEE7",
    "other": "60788F",
}


def export_xlsx(data: dict[str, Any], out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Métrés"
    headers = ["Ouvrage", "Catégorie", "Matériau", "Nombre", "Longueur (m)", "Surface murs (m²)", "Aire (m²)", "Volume (m³)", "Longueur dessin", "Aire dessin²"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="17324A")
        c.alignment = Alignment(horizontal="center")
    for q in data.get("quantities", []):
        ws.append([
            q.get("label"), q.get("category"), q.get("material"), q.get("count"), q.get("length_m"), q.get("surface_m2"),
            q.get("area_m2"), q.get("volume_m3"), q.get("linear_drawing_units"), q.get("area_drawing_units2")
        ])
        row = ws.max_row
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=CATEGORY_COLOR.get(q.get("category"), "60788F"))
    widths = [28, 18, 24, 10, 16, 18, 14, 14, 18, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    cfg = wb.create_sheet("Paramètres")
    cfg.append(["Paramètre", "Valeur"])
    cfg["A1"].font = cfg["B1"].font = Font(bold=True, color="FFFFFF")
    cfg["A1"].fill = cfg["B1"].fill = PatternFill("solid", fgColor="17324A")
    settings = data.get("settings", {})
    cfg.append(["Fichier", data.get("file")])
    cfg.append(["Format", data.get("kind")])
    cfg.append(["Unité", settings.get("unit")])
    cfg.append(["Facteur vers m", settings.get("unit_to_m")])
    for cat, vals in settings.get("category_defaults", {}).items():
        for key, value in vals.items():
            cfg.append([f"{cat}.{key}", value])
    cfg.column_dimensions["A"].width = 36
    cfg.column_dimensions["B"].width = 32

    layers = wb.create_sheet("Calques")
    layers.append(["Calque", "Nombre d'entités", "Catégorie"])
    for c in layers[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="17324A")
    for layer, count in data.get("summary", {}).get("layers", {}).items():
        layers.append([layer, count, settings.get("layer_map", {}).get(layer, "other")])
    layers.column_dimensions["A"].width = 42
    layers.column_dimensions["B"].width = 18
    layers.column_dimensions["C"].width = 22
    wb.save(out)


def export_csv(data: dict[str, Any], out: Path) -> None:
    fields = ["label", "category", "material", "count", "length_m", "surface_m2", "area_m2", "volume_m3", "linear_drawing_units", "area_drawing_units2"]
    with out.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, delimiter=";")
        w.writeheader()
        for row in data.get("quantities", []):
            w.writerow({k: row.get(k, "") for k in fields})


def _segments(entity: dict[str, Any]) -> list[tuple[tuple[float, float], tuple[float, float]]]:
    if entity.get("type") == "line":
        return [(tuple(entity["a"]), tuple(entity["b"]))]
    if entity.get("type") == "polyline":
        pts = [tuple(p) for p in entity.get("points", [])]
        segs = list(zip(pts, pts[1:]))
        if entity.get("closed") and len(pts) > 2:
            segs.append((pts[-1], pts[0]))
        return segs
    return []


def export_obj(data: dict[str, Any], out: Path) -> None:
    """Simple local OBJ export: each linear building element becomes a rectangular prism."""
    factor = data.get("settings", {}).get("unit_to_m") or 1.0
    defaults = data.get("settings", {}).get("category_defaults", {})
    vertices: list[tuple[float, float, float]] = []
    faces: list[tuple[int, int, int, int]] = []

    def add_box(a: tuple[float, float], b: tuple[float, float], width: float, z0: float, height: float) -> None:
        ax, ay = a[0] * factor, a[1] * factor
        bx, by = b[0] * factor, b[1] * factor
        dx, dy = bx - ax, by - ay
        length = math.hypot(dx, dy)
        if length < 1e-9:
            return
        nx, ny = -dy / length * width / 2, dx / length * width / 2
        base = len(vertices) + 1
        vertices.extend([
            (ax + nx, ay + ny, z0), (ax - nx, ay - ny, z0), (bx - nx, by - ny, z0), (bx + nx, by + ny, z0),
            (ax + nx, ay + ny, z0 + height), (ax - nx, ay - ny, z0 + height), (bx - nx, by - ny, z0 + height), (bx + nx, by + ny, z0 + height),
        ])
        faces.extend([
            (base, base + 1, base + 2, base + 3), (base + 4, base + 7, base + 6, base + 5),
            (base, base + 4, base + 5, base + 1), (base + 1, base + 5, base + 6, base + 2),
            (base + 2, base + 6, base + 7, base + 3), (base + 3, base + 7, base + 4, base),
        ])

    for ent in data.get("preview", {}).get("entities", []):
        cat = ent.get("cat", "other")
        if cat in {"other", "opening"}:
            continue
        meta = defaults.get(cat, {})
        width = float(meta.get("thickness") or meta.get("width") or 0.20)
        height = float(meta.get("height") or meta.get("thickness") or 0.20)
        z0 = 0.0
        if cat == "beam":
            z0 = max(0.0, float(defaults.get("wall", {}).get("height", 2.8)) - height)
        for a, b in _segments(ent):
            add_box(a, b, width, z0, height)

    with out.open("w", encoding="utf-8") as f:
        f.write("# BTP 3D Studio OBJ export\n")
        f.write("o BTP3D_Model\n")
        for x, y, z in vertices:
            f.write(f"v {x:.6f} {y:.6f} {z:.6f}\n")
        for face in faces:
            f.write("f " + " ".join(map(str, face)) + "\n")


def export_report_html(data: dict[str, Any], out: Path) -> None:
    qs = data.get("quantities", [])
    vol = sum(float(x.get("volume_m3") or 0) for x in qs)
    surf = sum(float(x.get("surface_m2") or 0) for x in qs)
    rows = "".join(
        f"<tr><td>{html.escape(str(q.get('label','')))}</td><td>{html.escape(str(q.get('material','')))}</td><td>{q.get('count','')}</td>"
        f"<td>{q.get('length_m','')}</td><td>{q.get('surface_m2','')}</td><td>{q.get('area_m2','')}</td><td>{q.get('volume_m3','')}</td></tr>"
        for q in qs
    )
    warnings = "".join(f"<li>{html.escape(str(w))}</li>" for w in data.get("warnings", []))
    payload = json.dumps(data, ensure_ascii=False)
    out.write_text(f"""<!doctype html><html lang='fr'><head><meta charset='utf-8'><title>Rapport BTP 3D Studio</title>
<style>body{{font-family:Arial,sans-serif;margin:34px;color:#183044}}h1{{margin-bottom:4px}}.muted{{color:#61778c}}.cards{{display:flex;gap:14px;margin:22px 0}}.card{{border:1px solid #d9e3eb;border-radius:12px;padding:14px 18px;min-width:150px}}table{{border-collapse:collapse;width:100%;font-size:13px}}th,td{{border-bottom:1px solid #dfe7ed;padding:9px;text-align:left}}th{{background:#17324a;color:white}}.foot{{margin-top:30px;font-size:11px;color:#71879a}}@media print{{button{{display:none}}}}</style></head><body>
<h1>BTP 3D Studio — Rapport de pré-métré</h1><div class='muted'>{html.escape(str(data.get('file','')))} · {html.escape(str(data.get('kind','')))}</div>
<div class='cards'><div class='card'><b>{vol:.2f} m³</b><div class='muted'>Volume pré-métré</div></div><div class='card'><b>{surf:.2f} m²</b><div class='muted'>Surface murs</div></div><div class='card'><b>{len(qs)}</b><div class='muted'>Familles d'ouvrages</div></div></div>
<h2>Métrés</h2><table><thead><tr><th>Ouvrage</th><th>Matériau</th><th>Nb</th><th>Longueur m</th><th>Surface m²</th><th>Aire m²</th><th>Volume m³</th></tr></thead><tbody>{rows}</tbody></table>
<h2>Alertes / hypothèses</h2><ul>{warnings or '<li>Aucune alerte.</li>'}</ul>
<p class='foot'>Pré-métré automatique à vérifier avant utilisation contractuelle ou financière. Généré localement par BTP 3D Studio v1.0.</p>
<script type='application/json' id='btp3d-data'>{html.escape(payload)}</script></body></html>""", encoding="utf-8")

_IFC64 = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz_$"

def _ifc_guid(seed: int) -> str:
    # Compact 128-bit-like identifier using IFC's 64-character alphabet.
    n = (0x9E3779B97F4A7C15 * (seed + 1)) & ((1 << 128) - 1)
    n = (n << 64) | ((0xD1B54A32D192ED03 * (seed + 7)) & ((1 << 64) - 1))
    chars = []
    for _ in range(22):
        chars.append(_IFC64[n & 63]); n >>= 6
    return ''.join(reversed(chars))


def export_ifc(data: dict[str, Any], out: Path) -> None:
    """Write a compact IFC4 model with extruded linear elements.

    The exporter is dependency-free and intentionally conservative: linear walls,
    foundations and beams are represented as rectangular extrusions. It is meant
    for coordination/visual checking, not fabrication detailing.
    """
    factor = float(data.get("settings", {}).get("unit_to_m") or 1.0)
    defaults = data.get("settings", {}).get("category_defaults", {})
    lines: list[str] = []
    next_id = 1
    def add(expr: str) -> int:
        nonlocal next_id
        i = next_id; next_id += 1; lines.append(f"#{i}={expr};"); return i
    def pt(x=0.0,y=0.0,z=0.0): return add(f"IFCCARTESIANPOINT(({x:.6f},{y:.6f},{z:.6f}))")
    def direction(x,y,z): return add(f"IFCDIRECTION(({x:.6f},{y:.6f},{z:.6f}))")
    def axis3(x=0,y=0,z=0,angle=0.0):
        p=pt(x,y,z); dz=direction(0,0,1); dx=direction(math.cos(angle),math.sin(angle),0); return add(f"IFCAXIS2PLACEMENT3D(#{p},#{dz},#{dx})")
    def local(x=0,y=0,z=0,angle=0.0,parent=None):
        ax=axis3(x,y,z,angle); return add(f"IFCLOCALPLACEMENT({('#'+str(parent)) if parent else '$'},#{ax})")

    # Ownership + units + representation context
    person=add("IFCPERSON($,$,'BTP 3D Studio',$,$,$,$,$)")
    org=add("IFCORGANIZATION($,'BTP 3D Studio',$,$,$)")
    pao=add(f"IFCPERSONANDORGANIZATION(#{person},#{org},$)")
    app=add(f"IFCAPPLICATION(#{org},'1.0','BTP 3D Studio','BTP3D')")
    owner=add(f"IFCOWNERHISTORY(#{pao},#{app},$,.ADDED.,$,$,$,0)")
    wcs=axis3()
    context=add(f"IFCGEOMETRICREPRESENTATIONCONTEXT($,'Model',3,1.E-05,#{wcs},$)")
    ulen=add("IFCSIUNIT(*,.LENGTHUNIT.,$,.METRE.)")
    uarea=add("IFCSIUNIT(*,.AREAUNIT.,$,.SQUARE_METRE.)")
    uvol=add("IFCSIUNIT(*,.VOLUMEUNIT.,$,.CUBIC_METRE.)")
    units=add(f"IFCUNITASSIGNMENT((#{ulen},#{uarea},#{uvol}))")
    project=add(f"IFCPROJECT('{_ifc_guid(1)}',#{owner},'BTP3D Project',$,$,$,$,(#{context}),#{units})")
    site_place=local(); site=add(f"IFCSITE('{_ifc_guid(2)}',#{owner},'Site',$,$,#{site_place},$,$,.ELEMENT.,$,$,$,$,$)")
    bld_place=local(parent=site_place); bld=add(f"IFCBUILDING('{_ifc_guid(3)}',#{owner},'Bâtiment',$,$,#{bld_place},$,$,.ELEMENT.,$,$,$)")
    sty_place=local(parent=bld_place); sty=add(f"IFCBUILDINGSTOREY('{_ifc_guid(4)}',#{owner},'Niveau 0',$,$,#{sty_place},$,$,.ELEMENT.,0.)")
    add(f"IFCRELAGGREGATES('{_ifc_guid(5)}',#{owner},$,$,#{project},(#{site}))")
    add(f"IFCRELAGGREGATES('{_ifc_guid(6)}',#{owner},$,$,#{site},(#{bld}))")
    add(f"IFCRELAGGREGATES('{_ifc_guid(7)}',#{owner},$,$,#{bld},(#{sty}))")

    element_ids=[]; counter=20
    for ent in data.get("preview", {}).get("entities", []):
        cat=ent.get("cat","other")
        if cat not in {"wall","foundation","beam","column"}: continue
        meta=defaults.get(cat,{})
        width=float(meta.get("thickness") or meta.get("width") or 0.20)
        height=float(meta.get("height") or 0.25)
        z0=0.0
        if cat=="beam": z0=max(0.0,float(defaults.get("wall",{}).get("height",2.8))-height)
        for a,b in _segments(ent):
            ax,ay=a[0]*factor,a[1]*factor; bx,by=b[0]*factor,b[1]*factor
            dx,dy=bx-ax,by-ay; length=math.hypot(dx,dy)
            if length<1e-6: continue
            angle=math.atan2(dy,dx)
            place=local(ax,ay,z0,angle,parent=sty_place)
            prof_place=add(f"IFCAXIS2PLACEMENT2D(#{pt(0,0,0)},$)")
            prof=add(f"IFCRECTANGLEPROFILEDEF(.AREA.,$,#{prof_place},{length:.6f},{width:.6f})")
            solid_pos=axis3(length/2,0,0,0)
            extrude_dir=direction(0,0,1)
            solid=add(f"IFCEXTRUDEDAREASOLID(#{prof},#{solid_pos},#{extrude_dir},{height:.6f})")
            rep=add(f"IFCSHAPEREPRESENTATION(#{context},'Body','SweptSolid',(#{solid}))")
            shape=add(f"IFCPRODUCTDEFINITIONSHAPE($,$,(#{rep}))")
            counter+=1
            name=defaults.get(cat,{}).get('label',cat).replace("'","''")
            if cat=='wall': expr=f"IFCWALL('{_ifc_guid(counter)}',#{owner},'{name}',$,$,#{place},#{shape},$,.NOTDEFINED.)"
            elif cat=='beam': expr=f"IFCBEAM('{_ifc_guid(counter)}',#{owner},'{name}',$,$,#{place},#{shape},$,.NOTDEFINED.)"
            elif cat=='column': expr=f"IFCCOLUMN('{_ifc_guid(counter)}',#{owner},'{name}',$,$,#{place},#{shape},$,.NOTDEFINED.)"
            else: expr=f"IFCFOOTING('{_ifc_guid(counter)}',#{owner},'{name}',$,$,#{place},#{shape},$,.STRIP_FOOTING.)"
            element_ids.append(add(expr))
    if element_ids:
        ids=','.join('#'+str(i) for i in element_ids)
        add(f"IFCRELCONTAINEDINSPATIALSTRUCTURE('{_ifc_guid(999)}',#{owner},$,$,({ids}),#{sty})")

    header="""ISO-10303-21;\nHEADER;\nFILE_DESCRIPTION(('ViewDefinition [CoordinationView]'),'2;1');\nFILE_NAME('BTP3D.ifc','2026-09-26T00:00:00',('BTP 3D Studio'),('OpenAI'),'BTP 3D Studio','BTP 3D Studio','');\nFILE_SCHEMA(('IFC4'));\nENDSEC;\nDATA;\n"""
    out.write_text(header+'\n'.join(lines)+"\nENDSEC;\nEND-ISO-10303-21;\n",encoding='utf-8')
